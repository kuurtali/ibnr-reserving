r"""IBNR rezerv hesabi - ceyreklik incurred chain ladder.

Girdi: hasar kaydi tablosu.  Cikti: gelisim ucgeni, link ratio, kuyruk
faktoru, ultimate ve IBNR; grafikli bir Excel raporu olarak. Faktor
penceresi senaryolari (all / 5y / 3y) ve geriye donuk test dahil.

VERI POLITIKASI - bu depo gercek portfoy verisi icermez. Varsayilan girdi
data/ornek_hasar.xlsx (sentetik, bkz. data/uret.py). Gercek veri yolu
ortam degiskeninden verilir:  set IBNR_GIRIS_DOSYASI=D:\yol\hasar.xlsx
"""

import os
import re
import sys
import json
import subprocess

import numpy as np
import pandas as pd
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter, column_index_from_string
from openpyxl.formatting.rule import ColorScaleRule
from openpyxl.worksheet.table import Table, TableStyleInfo
from openpyxl.chart import BarChart, PieChart, LineChart, Reference
from openpyxl.chart.data_source import StrRef, AxDataSource

pd.options.display.float_format = '{:,.2f}'.format

INTERAKTIF = True
AYAR_DOSYASI = os.path.join(os.path.expanduser("~"), ".ibnr_ayarlar.json")
DONEM_DESENI = r'^\d{4}Q[1-4]$'
ZORUNLU_KOLONLAR = ['HasarDonem', 'GelDonemHT', 'HazAltBrnKod', 'Bedeni',
                    'Havuz', 'HasarDurum', 'HasarTutar']
PENCERE_SECENEKLERI = ['all', 'last_5y', 'last_3y']
FARK_KOLONU = "all'a Göre Fark %"
SEGMENTLER = [('maddi', 'Maddî', 0), ('bedeni', 'Bedeni', 1)]

VARSAYILAN_AYAR = {
    'giris': os.environ.get(
        "IBNR_GIRIS_DOSYASI",
        os.path.join(os.path.dirname(os.path.abspath(__file__)),
                     "data", "ornek_hasar.xlsx")),
    'degerleme': None,
    'brans': 715,
    'havuz': 0,
    'ceyrek': 40,
    'pencere': 'all',
    'tail': 1.00,
    'tail_otomatik': False,
    'sovtaj': 1,
    'senaryolar': True,
    'geriye_test': True,
}


def donem_kodu(ceyrek):
    ceyrek = str(ceyrek).strip().upper()
    if not re.match(DONEM_DESENI, ceyrek):
        raise ValueError(f"Geçersiz dönem: {ceyrek!r}. Beklenen biçim: 2026Q2")
    return int(ceyrek[:4]) * 4 + int(ceyrek[-1])


def kod_donem(kod):
    kod = int(kod)
    yil = kod // 4 if kod % 4 != 0 else (kod // 4) - 1
    q = kod % 4 if kod % 4 != 0 else 4
    return f"{yil}Q{q}"


def tail_dogrula(x):
    deger = float(str(x).replace(',', '.'))
    if deger <= 0:
        raise ValueError("Kuyruk faktörü pozitif bir sayı olmalı")
    return deger


def isaret_dogrula(x):
    deger = tam_sayi(x)
    if deger not in (1, -1):
        raise ValueError("Sovtaj işareti 1 veya -1 olmalı")
    return deger


def pencere_dogrula(x):
    deger = str(x).strip().lower()
    if deger not in PENCERE_SECENEKLERI:
        raise ValueError("Faktör penceresi: all, last_5y veya last_3y")
    return deger


def sayi_metni(deger):
    try:
        sayi = float(deger)
    except (TypeError, ValueError):
        return str(deger)
    return str(int(sayi)) if float(sayi).is_integer() else str(sayi)


def tam_sayi(x):
    metin = str(x).strip().replace(',', '.')
    if metin == "":
        raise ValueError("Boş bırakılamaz")
    return int(float(metin))


def tl(x):
    return f"{x:,.0f} TL"


def ayarlari_oku():
    ayar = dict(VARSAYILAN_AYAR)
    try:
        with open(AYAR_DOSYASI, encoding='utf-8') as dosya:
            kayitli = json.load(dosya)
        for anahtar in VARSAYILAN_AYAR:
            if anahtar in kayitli:
                ayar[anahtar] = kayitli[anahtar]
    except (OSError, ValueError):
        pass
    return ayar


def ayarlari_yaz(ayar):
    kayit = {anahtar: ayar.get(anahtar) for anahtar in VARSAYILAN_AYAR}
    try:
        with open(AYAR_DOSYASI, 'w', encoding='utf-8') as dosya:
            json.dump(kayit, dosya, ensure_ascii=False, indent=2)
    except OSError:
        pass


def dosya_ac(yol):
    try:
        if os.name == 'nt':
            os.startfile(yol)
        elif sys.platform == 'darwin':
            subprocess.run(['open', yol], check=False)
        else:
            subprocess.run(['xdg-open', yol], check=False)
    except Exception as hata:
        print(f"Dosya açılamadı: {hata}")


def veri_oku(yol):
    if not os.path.isfile(yol):
        raise SystemExit(f"Dosya bulunamadı: {yol}")
    ham = pd.read_excel(yol)
    eksik = [k for k in ZORUNLU_KOLONLAR if k not in ham.columns]
    if eksik:
        raise SystemExit(f"Dosyada eksik kolon(lar): {', '.join(eksik)}")

    kalite = [("Dosyadaki toplam satır", len(ham), "")]
    donem_maske = ham['HasarDonem'].astype(str).str.strip().str.match(DONEM_DESENI, case=False)
    kalite.append(("Geçersiz HasarDonem biçimi", int((~donem_maske).sum()),
                   "2026Q2 biçiminde olmayan satırlar hesaba alınmadı"))
    df = ham[donem_maske].copy()
    if df.empty:
        raise SystemExit(f"'{yol}' içinde 2026Q2 biçiminde geçerli HasarDonem bulunamadı.")

    tutar = pd.to_numeric(df['HasarTutar'], errors='coerce')
    kalite.append(("HasarTutar sayı değil / boş", int(tutar.isna().sum()), "0 olarak alındı"))
    kalite.append(("HasarTutar negatif", int((tutar < 0).sum()), "olduğu gibi kullanıldı"))

    gel = pd.to_numeric(df['GelDonemHT'], errors='coerce')
    kalite.append(("GelDonemHT boş", int(gel.isna().sum()), "hesaba alınmadı"))
    kalite.append(("GelDonemHT negatif", int((gel < 0).sum()), "hesaba alınmadı"))

    durumlar = df['HasarDurum'].astype(str).str.strip().value_counts()
    beklenen = ['Odenen', 'Muallak', 'SovtajRucu']
    for durum, adet in durumlar.items():
        kalite.append((f"HasarDurum = {durum}", int(adet),
                       "" if durum in beklenen else "BEKLENMEYEN DEĞER — hesaba alınmadı"))
    for kolon in ('HazAltBrnKod', 'Havuz', 'Bedeni', 'GelDonemHT'):
        df[kolon] = pd.to_numeric(df[kolon], errors='coerce')
    df['HasarDurum'] = df['HasarDurum'].astype(str).str.strip()
    kodlar = df['HasarDonem'].apply(donem_kodu)
    meta = {
        'yol': yol,
        'veri_ilk': int(kodlar.min()),
        'veri_son': int(kodlar.max()),
        'brans_secenekleri': [sayi_metni(k) for k in sorted(df['HazAltBrnKod'].dropna().unique())],
        'havuz_secenekleri': [sayi_metni(k) for k in sorted(df['Havuz'].dropna().unique())],
        'kalite': kalite,
        'satir_sayisi': len(df),
    }
    return df, meta


def tail_tahmin(faktorler, ileri_adim=20, ust_sinir=1.5):
    noktalar = [(j, f - 1.0) for j, f in enumerate(faktorler) if f > 1.0005]
    if len(noktalar) < 3:
        return 1.0, "artan faktör yok"
    noktalar = noktalar[-8:]
    x = np.array([j for j, _ in noktalar], dtype=float)
    y = np.log(np.array([d for _, d in noktalar], dtype=float))
    egim, kesisim = np.polyfit(x, y, 1)
    if egim >= 0:
        return 1.0, "faktörler azalmıyor"
    son_j = int(x[-1])
    tail = 1.0
    for j in range(son_j + 1, son_j + 1 + ileri_adim):
        tail *= 1.0 + float(np.exp(kesisim + egim * j))
    tail = min(max(tail, 1.0), ust_sinir)
    return round(tail, 6), f"{len(noktalar)} noktaya üstel uyum"


def ayari_dogrula(ayar, meta):
    uyarilar = []
    kod = donem_kodu(ayar['degerleme'])
    if kod <= meta['veri_ilk']:
        raise ValueError(f"Değerleme dönemi en erken {kod_donem(meta['veri_ilk'] + 1)} olabilir; "
                         f"veri {kod_donem(meta['veri_ilk'])} döneminde başlıyor.")
    ayar['degerleme'] = kod_donem(kod)
    ayar['degerleme_kodu'] = kod

    ceyrek = tam_sayi(ayar['ceyrek'])
    if ceyrek < 2:
        raise ValueError("Üçgen en az 2 çeyrek olmalı.")
    azami = kod - meta['veri_ilk'] + 1
    if ceyrek > azami:
        uyarilar.append(f"{ceyrek} çeyrek istendi; veri {azami} çeyrek destekliyor, {azami} kullanıldı.")
        ceyrek = azami
    ayar['ceyrek'] = ceyrek

    if kod > meta['veri_son']:
        uyarilar.append(f"Değerleme dönemi verinin {kod - meta['veri_son']} çeyrek ilerisinde; "
                        f"{kod_donem(meta['veri_son'] + 1)} ve sonrası boş kalır, IBNR üretmez.")
    return uyarilar


def segment_hesapla(df, bedeni_flag, ayar, meta, faktor_override=None):
    periyot = tam_sayi(ayar['ceyrek'])
    if periyot < 2:
        raise ValueError("Üçgen en az 2 çeyrek olmalı.")
    son_indeks = periyot - 1
    baslangic = int(ayar['degerleme_kodu']) - periyot + 1
    diyagonal = min(meta['veri_son'], int(ayar['degerleme_kodu'])) - baslangic
    if diyagonal < 0:
        raise ValueError("Seçilen dönem aralığında hiç veri yok.")

    df_sub = df[
        (df['HazAltBrnKod'] == float(ayar['brans'])) &
        (df['Bedeni'] == float(bedeni_flag)) &
        (df['Havuz'] == float(ayar['havuz'])) &
        (df['GelDonemHT'].notna())
    ].copy()

    secili_satir = len(df_sub)
    df_sub = df_sub[df_sub['GelDonemHT'] >= 0].copy()
    df_sub['Hasar_Idx'] = df_sub['HasarDonem'].apply(donem_kodu) - baslangic
    df_sub['Gel_Idx'] = df_sub['GelDonemHT'].astype(int)
    donem_disi = int(((df_sub['Hasar_Idx'] < 0) | (df_sub['Hasar_Idx'] > son_indeks)).sum())
    diyagonal_disi = int((df_sub['Hasar_Idx'] + df_sub['Gel_Idx'] > diyagonal).sum())
    df_sub = df_sub[
        (df_sub['Hasar_Idx'] >= 0) &
        (df_sub['Hasar_Idx'] <= son_indeks) &
        (df_sub['Gel_Idx'] >= 0) &
        (df_sub['Hasar_Idx'] + df_sub['Gel_Idx'] <= diyagonal)
    ].copy()

    df_sub['Tutar'] = pd.to_numeric(df_sub['HasarTutar'], errors='coerce').fillna(0.0)
    df_sub.loc[df_sub['HasarDurum'] == 'SovtajRucu', 'Tutar'] *= tam_sayi(ayar['sovtaj'])

    odenen = df_sub[df_sub['HasarDurum'].isin(['Odenen', 'SovtajRucu'])].pivot_table(
        index='Hasar_Idx', columns='Gel_Idx', values='Tutar', aggfunc='sum'
    ).reindex(index=range(periyot), columns=range(periyot), fill_value=0)

    muallak = df_sub[df_sub['HasarDurum'] == 'Muallak'].pivot_table(
        index='Hasar_Idx', columns='Gel_Idx', values='Tutar', aggfunc='sum'
    ).reindex(index=range(periyot), columns=range(periyot), fill_value=0)

    ucgen = odenen.cumsum(axis=1) + muallak
    maske = np.add.outer(np.arange(periyot), np.arange(periyot)) > diyagonal
    ucgen = ucgen.where(~maske, 0.0)

    hesaplanan = []
    for j in range(son_indeks):
        gecerli = ucgen[(ucgen[j] > 0) & (ucgen[j + 1] > 0)]
        if ayar['pencere'] == 'last_3y':
            gecerli = gecerli.tail(12)
        elif ayar['pencere'] == 'last_5y':
            gecerli = gecerli.tail(20)
        payda = gecerli[j].sum()
        hesaplanan.append(gecerli[j + 1].sum() / payda if payda != 0 else 1.0)

    if ayar.get('tail_otomatik'):
        tail, tail_not = tail_tahmin(hesaplanan)
    else:
        tail, tail_not = float(ayar['tail']), "elle girildi"

    faktorler = list(hesaplanan)
    if faktor_override:
        faktorler = [float(x) for x in faktor_override[:son_indeks]]
        if len(faktor_override) > son_indeks:
            tail = float(faktor_override[son_indeks])
            tail_not = "elle düzeltildi"
        while len(faktorler) < son_indeks:
            faktorler.append(1.0)

    projekte = ucgen.copy()
    for r in range(periyot):
        for c in range(1, periyot):
            if r + c > diyagonal:
                projekte.iloc[r, c] = projekte.iloc[r, c - 1] * faktorler[c - 1]
    projekte[son_indeks] = projekte[son_indeks] * tail

    etiketler = [kod_donem(baslangic + i) for i in range(periyot)]
    detay = pd.DataFrame(index=etiketler)
    detay['Mevcut_Kumulatif'] = [
        ucgen.iloc[i, diyagonal - i] if 0 <= diyagonal - i < periyot else 0.0
        for i in range(periyot)
    ]
    detay['Nihai_Hasar_Ultimate'] = projekte[son_indeks].values
    detay['IBNR'] = detay['Nihai_Hasar_Ultimate'] - detay['Mevcut_Kumulatif']

    return {
        'sayimlar': {
            'secili': secili_satir,
            'donem_disi': donem_disi,
            'diyagonal_disi': diyagonal_disi,
            'kullanilan': len(df_sub),
        },
        'ucgen': ucgen,
        'projekte': projekte,
        'detay': detay,
        'faktorler': faktorler,
        'hesaplanan_faktorler': hesaplanan,
        'tail': tail,
        'tail_not': tail_not,
        'kumulatif': float(detay['Mevcut_Kumulatif'].sum()),
        'ultimate': float(detay['Nihai_Hasar_Ultimate'].sum()),
        'ibnr': float(detay['IBNR'].sum()),
        'etiketler': etiketler,
        'periyot': periyot,
        'son_indeks': son_indeks,
        'diyagonal': diyagonal,
        'baslangic': baslangic,
    }


def hesapla(df, ayar, meta, override=None, durum=None):
    if durum:
        durum("Üçgenler kuruluyor...")
    sonuc = {}
    for anahtar, _ad, bayrak in SEGMENTLER:
        sonuc[anahtar] = segment_hesapla(
            df, bayrak, ayar, meta,
            faktor_override=(override or {}).get(anahtar)
        )

    maddi, bedeni = sonuc['maddi'], sonuc['bedeni']
    sonuc['ozet'] = pd.DataFrame({
        'Mevcut Kümülatif (Gerçekleşen)': [maddi['kumulatif'], bedeni['kumulatif'],
                                           maddi['kumulatif'] + bedeni['kumulatif']],
        'Nihai Hasar (Ultimate)': [maddi['ultimate'], bedeni['ultimate'],
                                   maddi['ultimate'] + bedeni['ultimate']],
        'Hesaplanan IBNR': [maddi['ibnr'], bedeni['ibnr'], maddi['ibnr'] + bedeni['ibnr']],
    }, index=['Maddî (Bedeni=0)', 'Bedeni (Bedeni=1)', 'GENEL TOPLAM'])

    sonuc['toplam_kumulatif'] = maddi['kumulatif'] + bedeni['kumulatif']
    sonuc['toplam_ultimate'] = maddi['ultimate'] + bedeni['ultimate']
    sonuc['toplam_ibnr'] = maddi['ibnr'] + bedeni['ibnr']
    sonuc['etiketler'] = maddi['etiketler']
    sonuc['periyot'] = maddi['periyot']
    sonuc['son_indeks'] = maddi['son_indeks']
    sonuc['diyagonal'] = maddi['diyagonal']
    sonuc['veri_son_donemi'] = kod_donem(maddi['baslangic'] + maddi['diyagonal'])
    sonuc['baslangic_donemi'] = maddi['etiketler'][0]

    if ayar.get('senaryolar') and not override:
        if durum:
            durum("Faktör senaryoları karşılaştırılıyor...")
        satirlar = []
        for pencere in PENCERE_SECENEKLERI:
            senaryo_ayar = dict(ayar)
            senaryo_ayar['pencere'] = pencere
            m = segment_hesapla(df, 0, senaryo_ayar, meta)
            b = segment_hesapla(df, 1, senaryo_ayar, meta)
            satirlar.append({
                'Faktör Penceresi': pencere,
                'Maddî Ultimate': m['ultimate'],
                'Bedeni Ultimate': b['ultimate'],
                'Toplam Ultimate': m['ultimate'] + b['ultimate'],
                'Toplam IBNR': m['ibnr'] + b['ibnr'],
                'Kuyruk Faktörü': m['tail'],
            })
        senaryo = pd.DataFrame(satirlar)
        temel = senaryo.loc[senaryo['Faktör Penceresi'] == 'all', 'Toplam Ultimate'].iloc[0]
        senaryo[FARK_KOLONU] = (senaryo['Toplam Ultimate'] / temel - 1) * 100 if temel else 0.0
        sonuc['senaryolar'] = senaryo
    else:
        sonuc['senaryolar'] = None

    if ayar.get('geriye_test') and not override:
        if durum:
            durum("Geriye dönük test yapılıyor...")
        sonuc['geriye_test'] = geriye_donuk_test(df, ayar, meta, sonuc)
    else:
        sonuc['geriye_test'] = None

    return sonuc


def link_ratio_tablosu(segment):
    ucgen = segment['ucgen']
    periyot, son_indeks = segment['periyot'], segment['son_indeks']
    diyagonal = segment['diyagonal']
    sutunlar = [f"{j} -> {j+1}" for j in range(son_indeks)]
    oranlar = pd.DataFrame(index=segment['etiketler'], columns=sutunlar, dtype=float)
    for r in range(periyot):
        for j in range(son_indeks):
            if r + j + 1 <= diyagonal and ucgen.iloc[r, j] > 0 and ucgen.iloc[r, j + 1] > 0:
                oranlar.iloc[r, j] = ucgen.iloc[r, j + 1] / ucgen.iloc[r, j]

    ozet = pd.DataFrame(index=['Hacim ağırlıklı', 'Basit ortalama', 'Medyan', 'Adet',
                               'Kullanılan faktör'], columns=sutunlar, dtype=float)
    for j, sutun in enumerate(sutunlar):
        gecerli = ucgen[(ucgen[j] > 0) & (ucgen[j + 1] > 0)]
        payda = gecerli[j].sum()
        ozet.loc['Hacim ağırlıklı', sutun] = gecerli[j + 1].sum() / payda if payda else np.nan
        ozet.loc['Basit ortalama', sutun] = oranlar[sutun].mean()
        ozet.loc['Medyan', sutun] = oranlar[sutun].median()
        ozet.loc['Adet', sutun] = oranlar[sutun].count()
        ozet.loc['Kullanılan faktör', sutun] = segment['faktorler'][j]
    return pd.concat([oranlar, ozet])


def geriye_donuk_test(df, ayar, meta, simdiki, gecikmeler=(4, 8, 12)):
    satirlar = []
    simdiki_ult = (simdiki['maddi']['detay']['Nihai_Hasar_Ultimate'] +
                   simdiki['bedeni']['detay']['Nihai_Hasar_Ultimate'])
    for gecikme in gecikmeler:
        eski_kod = int(ayar['degerleme_kodu']) - gecikme
        if eski_kod <= meta['veri_ilk'] + 1:
            continue
        eski_ayar = dict(ayar)
        eski_ayar['degerleme'] = kod_donem(eski_kod)
        eski_ayar['degerleme_kodu'] = eski_kod
        eski_ayar['ceyrek'] = min(int(ayar['ceyrek']), eski_kod - meta['veri_ilk'] + 1)
        eski_ayar['senaryolar'] = False
        if eski_ayar['ceyrek'] < 2:
            continue
        try:
            maddi = segment_hesapla(df, 0, eski_ayar, meta)
            bedeni = segment_hesapla(df, 1, eski_ayar, meta)
        except ValueError:
            continue
        eski_ult = (maddi['detay']['Nihai_Hasar_Ultimate'] +
                    bedeni['detay']['Nihai_Hasar_Ultimate'])
        ortak = [d for d in eski_ult.index if d in simdiki_ult.index]
        if not ortak:
            continue
        o_gun = float(eski_ult.loc[ortak].sum())
        bugun = float(simdiki_ult.loc[ortak].sum())
        satirlar.append({
            'Gecikme (çeyrek)': gecikme,
            'Değerleme Dönemi': eski_ayar['degerleme'],
            'Ortak Kaza Dönemi': len(ortak),
            'O Günkü Ultimate Tahmini': o_gun,
            'Bugünkü Ultimate Tahmini': bugun,
            'Fark (TL)': bugun - o_gun,
            'Sapma %': (o_gun / bugun - 1) * 100 if bugun else 0.0,
        })
    return pd.DataFrame(satirlar) if satirlar else None


def varsayim_tablosu(sonuc, ayar, meta):
    maddi, bedeni = sonuc['maddi'], sonuc['bedeni']
    ezilmis = []
    for anahtar, ad, _ in SEGMENTLER:
        segment = sonuc[anahtar]
        farklar = [f"{j} -> {j+1}"
                   for j, (kullanilan, hesaplanan) in enumerate(
                       zip(segment['faktorler'], segment['hesaplanan_faktorler']))
                   if abs(kullanilan - hesaplanan) > 1e-9]
        if farklar:
            ezilmis.append(f"{ad}: {', '.join(farklar)}")

    satirlar = [
        ("Çalıştırma zamanı", f"{pd.Timestamp.now():%d.%m.%Y %H:%M:%S}"),
        ("Kaynak dosya", meta['yol']),
        ("Geçerli hasar satırı", f"{meta['satir_sayisi']:,}"),
        ("Veri aralığı", f"{kod_donem(meta['veri_ilk'])} - {kod_donem(meta['veri_son'])}"),
        ("Değerleme dönemi", ayar['degerleme']),
        ("Üçgen aralığı", f"{sonuc['baslangic_donemi']} - {ayar['degerleme']} "
                          f"({sonuc['periyot']} çeyrek)"),
        ("Veri diyagonali", sonuc['veri_son_donemi']),
        ("Branş kodu", sayi_metni(ayar['brans'])),
        ("Havuz kodu", sayi_metni(ayar['havuz'])),
        ("Yöntem", "Chain Ladder (kümülatif ödenen + muallak)"),
        ("Faktör penceresi", ayar['pencere']),
        ("Faktör hesabı", "Hacim ağırlıklı; j ve j+1 sütunu pozitif olan satırlar"),
        ("Kuyruk faktörü (Maddî)", f"{maddi['tail']:.6f} ({maddi['tail_not']})"),
        ("Kuyruk faktörü (Bedeni)", f"{bedeni['tail']:.6f} ({bedeni['tail_not']})"),
        ("Sovtaj/rücu işareti", f"{tam_sayi(ayar['sovtaj']):+d}"),
        ("Elle düzeltilen faktör", "; ".join(ezilmis) if ezilmis else "yok"),
        ("Kullanılan satır (Maddî)", f"{maddi['sayimlar']['kullanilan']:,}"),
        ("Kullanılan satır (Bedeni)", f"{bedeni['sayimlar']['kullanilan']:,}"),
        ("Python", sys.version.split()[0]),
        ("pandas / openpyxl", f"{pd.__version__} / {openpyxl.__version__}"),
    ]
    return pd.DataFrame(satirlar, columns=['Varsayım', 'Değer'])


def kalite_tablosu(sonuc, ayar, meta):
    satirlar = [{'Kontrol': ad, 'Adet': adet, 'Not': aciklama} for ad, adet, aciklama in meta['kalite']]
    satirlar.append({'Kontrol': "-- Seçilen branş/havuz --", 'Adet': "", 'Not': ""})
    for anahtar, ad, _ in SEGMENTLER:
        sayim = sonuc[anahtar]['sayimlar']
        satirlar.extend([
            {'Kontrol': f"{ad}: filtreye uyan satır", 'Adet': sayim['secili'], 'Not': ""},
            {'Kontrol': f"{ad}: üçgen dönem aralığı dışında", 'Adet': sayim['donem_disi'],
             'Not': "kaza dönemi seçilen aralıkta değil"},
            {'Kontrol': f"{ad}: diyagonal sonrası", 'Adet': sayim['diyagonal_disi'],
             'Not': "değerleme tarihinden sonraki gelişim"},
            {'Kontrol': f"{ad}: hesaba giren satır", 'Adet': sayim['kullanilan'], 'Not': ""},
        ])
    return pd.DataFrame(satirlar)


RENK_LACIVERT = PatternFill(start_color="1F4E78", end_color="1F4E78", fill_type="solid")
RENK_TOPLAM = PatternFill(start_color="D9E1F2", end_color="D9E1F2", fill_type="solid")
RENK_MAVI = PatternFill(start_color="EAF1F8", end_color="EAF1F8", fill_type="solid")
RENK_KIRMIZI = PatternFill(start_color="FBEAEA", end_color="FBEAEA", fill_type="solid")
RENK_YESIL = PatternFill(start_color="EAF7EE", end_color="EAF7EE", fill_type="solid")
FONT_BASLIK = Font(name="Arial", size=11, bold=True, color="FFFFFF")
FONT_KALIN = Font(name="Arial", size=11, bold=True)
FONT_NORMAL = Font(name="Arial", size=11)
FONT_BANNER = Font(name="Arial", size=18, bold=True, color="FFFFFF")
FONT_ALT = Font(name="Arial", size=10, italic=True, color="808080")
FONT_KUTU_DEGER = Font(name="Arial", size=16, bold=True, color="1F4E78")
FONT_KUTU_BASLIK = Font(name="Arial", size=10, bold=True, color="595959")
KENAR_KALIN = Border(left=Side(style='medium', color='BFBFBF'), right=Side(style='medium', color='BFBFBF'),
                     top=Side(style='medium', color='BFBFBF'), bottom=Side(style='medium', color='BFBFBF'))
KENAR_INCE = Border(left=Side(style='thin', color='D9D9D9'), right=Side(style='thin', color='D9D9D9'),
                    top=Side(style='thin', color='D9D9D9'), bottom=Side(style='thin', color='D9D9D9'))


def kategori_ayarla(chart, ref):
    for seri in chart.series:
        seri.cat = AxDataSource(strRef=StrRef(f=str(ref)))


def bolum_basligi(ws, hucre, metin):
    ws[hucre].value = f"  {metin}"
    ws[hucre].font = Font(name="Arial", size=12, bold=True, color="1F4E78")
    ws[hucre].border = Border(left=Side(style='thick', color='1F4E78'))
    ws[hucre].alignment = Alignment(horizontal="left", vertical="center")


def baslik_satiri(ws):
    for hucre in ws[1]:
        hucre.fill = RENK_LACIVERT
        hucre.font = FONT_BASLIK
        hucre.alignment = Alignment(horizontal="center", vertical="center")


def uzun_ucgen(sonuc, anahtar, ad):
    segment = sonuc[anahtar]
    periyot = segment['periyot']
    diyagonal = segment['diyagonal']
    satirlar = []
    for r in range(periyot):
        for c in range(periyot):
            gozlenmis = (r + c <= diyagonal)
            satirlar.append({
                'Segment': ad,
                'KazaDonemi': segment['etiketler'][r],
                'KazaDonemi_Idx': r,
                'GelisimDonemi_Idx': c,
                'Deger': segment['ucgen'].iloc[r, c] if gozlenmis else segment['projekte'].iloc[r, c],
                'VeriTipi': 'Gerçekleşen' if gozlenmis else 'Projekte (IBNR)',
            })
    return pd.DataFrame(satirlar)


def excel_yaz(sonuc, ayar, meta, cikis, durum=None):
    if durum:
        durum("Excel hazırlanıyor...")

    periyot = sonuc['periyot']
    son_indeks = sonuc['son_indeks']
    etiketler = sonuc['etiketler']
    maddi, bedeni = sonuc['maddi'], sonuc['bedeni']

    ucgen_etiketli = {}
    for anahtar, ad, _ in SEGMENTLER:
        for tur in ('ucgen', 'projekte'):
            tablo = sonuc[anahtar][tur].copy()
            tablo.index = etiketler
            tablo.columns = [f"Gelişim_{c}" for c in range(periyot)]
            ucgen_etiketli[(anahtar, tur)] = tablo

    uzun_detay = pd.concat([
        maddi['detay'].reset_index().rename(columns={'index': 'KazaDonemi'}).assign(Segment='Maddî'),
        bedeni['detay'].reset_index().rename(columns={'index': 'KazaDonemi'}).assign(Segment='Bedeni'),
    ], ignore_index=True)[['Segment', 'KazaDonemi', 'Mevcut_Kumulatif', 'Nihai_Hasar_Ultimate', 'IBNR']]

    uzun_faktor = pd.DataFrame({
        'GelisimAdimi': list(range(son_indeks)) * 2,
        'Segment': ['Maddî'] * son_indeks + ['Bedeni'] * son_indeks,
        'Faktor': list(maddi['faktorler']) + list(bedeni['faktorler']),
    })

    uzun_ozet = sonuc['ozet'].reset_index().rename(columns={'index': 'Segment'})
    uzun_ozet = uzun_ozet[uzun_ozet['Segment'] != 'GENEL TOPLAM']
    uzun_tumu = pd.concat([uzun_ucgen(sonuc, 'maddi', 'Maddî'),
                           uzun_ucgen(sonuc, 'bedeni', 'Bedeni')], ignore_index=True)

    varsayimlar = varsayim_tablosu(sonuc, ayar, meta)
    kalite = kalite_tablosu(sonuc, ayar, meta)
    link_ratio = {anahtar: link_ratio_tablosu(sonuc[anahtar]) for anahtar, _, _ in SEGMENTLER}

    os.makedirs(os.path.dirname(cikis) or ".", exist_ok=True)
    with pd.ExcelWriter(cikis, engine='openpyxl') as writer:
        sonuc['ozet'].to_excel(writer, sheet_name='Özet')
        if sonuc['senaryolar'] is not None:
            sonuc['senaryolar'].to_excel(writer, sheet_name='Senaryolar', index=False)
        if sonuc.get('geriye_test') is not None:
            sonuc['geriye_test'].to_excel(writer, sheet_name='Geriye Dönük Test', index=False)
        varsayimlar.to_excel(writer, sheet_name='Varsayımlar', index=False)
        kalite.to_excel(writer, sheet_name='Veri Kalitesi', index=False)
        link_ratio['maddi'].to_excel(writer, sheet_name='Maddi Link Ratio')
        link_ratio['bedeni'].to_excel(writer, sheet_name='Bedeni Link Ratio')
        ucgen_etiketli[('maddi', 'ucgen')].to_excel(writer, sheet_name='Maddi Üçgen')
        ucgen_etiketli[('maddi', 'projekte')].to_excel(writer, sheet_name='Maddi Tamamlanmış Üçgen')
        ucgen_etiketli[('bedeni', 'ucgen')].to_excel(writer, sheet_name='Bedeni Üçgen')
        ucgen_etiketli[('bedeni', 'projekte')].to_excel(writer, sheet_name='Bedeni Tamamlanmış Üçgen')
        maddi['detay'].to_excel(writer, sheet_name='Maddi Çeyreklik Detay')
        bedeni['detay'].to_excel(writer, sheet_name='Bedeni Çeyreklik Detay')
        uzun_ozet.to_excel(writer, sheet_name='PBI_Ozet', index=False)
        uzun_detay.to_excel(writer, sheet_name='PBI_Detay', index=False)
        uzun_faktor.to_excel(writer, sheet_name='PBI_Faktor', index=False)
        uzun_tumu.to_excel(writer, sheet_name='PBI_Ucgen', index=False)

    wb = openpyxl.load_workbook(cikis)

    ws_ozet = wb['Özet']
    ws_ozet.sheet_view.showGridLines = False
    ws_ozet.sheet_properties.tabColor = "1F4E78"
    baslik_satiri(ws_ozet)
    for satir in ws_ozet.iter_rows(min_row=2, max_row=4, min_col=1, max_col=4):
        toplam_mu = (satir[0].value == 'GENEL TOPLAM')
        for hucre in satir:
            hucre.font = FONT_KALIN if toplam_mu else FONT_NORMAL
            hucre.border = KENAR_INCE
            if toplam_mu:
                hucre.fill = RENK_TOPLAM
            if isinstance(hucre.value, (int, float)):
                hucre.number_format = '#,##0.00 "TL"'

    ws_db = wb.create_sheet('Dashboard', 0)
    ws_db.sheet_view.showGridLines = False
    ws_db.sheet_properties.tabColor = "1F4E78"
    for sutun in "ABCDEFGHIJKLMNOPQRST":
        ws_db.column_dimensions[sutun].width = 16

    ws_db.merge_cells("A1:N1")
    ws_db.row_dimensions[1].height = 34
    ws_db["A1"].value = "  IBNR Raporu"
    ws_db["A1"].font = FONT_BANNER
    ws_db["A1"].fill = RENK_LACIVERT
    ws_db["A1"].alignment = Alignment(horizontal="left", vertical="center")

    ws_db.merge_cells("A2:N2")
    ws_db["A2"].value = (f"  Değerleme: {ayar['degerleme']}  |  Üçgen: {sonuc['baslangic_donemi']}–"
                         f"{ayar['degerleme']} ({periyot} çeyrek)  |  Veri diyagonali: {sonuc['veri_son_donemi']}"
                         f"  |  Branş: {ayar['brans']}  |  Havuz: {ayar['havuz']}  |  "
                         f"Faktör: {ayar['pencere']}  |  Kuyruk: {maddi['tail']:.4f} ({maddi['tail_not']})"
                         f"  |  Sovtaj: {int(ayar['sovtaj']):+d}  |  {pd.Timestamp.now():%d.%m.%Y %H:%M}")
    ws_db["A2"].font = FONT_ALT
    ws_db["A2"].alignment = Alignment(horizontal="left", vertical="center")

    kutular = [
        ("Mevcut Kümülatif (TL)", "B4", "A", "C", RENK_MAVI),
        ("IBNR (TL)", "D4", "E", "G", RENK_KIRMIZI),
        ("Ultimate (TL)", "C4", "I", "K", RENK_YESIL),
    ]
    for baslik, kaynak, ilk, son, renk in kutular:
        ws_db.merge_cells(f"{ilk}4:{son}4")
        ws_db[f"{ilk}4"].value = baslik
        ws_db[f"{ilk}4"].font = FONT_KUTU_BASLIK
        ws_db[f"{ilk}4"].alignment = Alignment(horizontal="center", vertical="center")
        ws_db.merge_cells(f"{ilk}5:{son}6")
        hucre = ws_db[f"{ilk}5"]
        hucre.value = f"='Özet'!{kaynak}"
        hucre.font = FONT_KUTU_DEGER
        hucre.number_format = '#,##0 "TL"'
        hucre.alignment = Alignment(horizontal="center", vertical="center", shrink_to_fit=True)
        hucre.fill = renk
        for satir_no in (5, 6):
            for sutun_no in range(column_index_from_string(ilk), column_index_from_string(son) + 1):
                ws_db.cell(row=satir_no, column=sutun_no).border = KENAR_KALIN

    bolum_basligi(ws_db, "A8", "Genel Karşılaştırma")

    grafik1 = BarChart()
    grafik1.type = "col"
    grafik1.style = 10
    grafik1.title = "Ultimate ve IBNR"
    grafik1.y_axis.title = "Tutar (TL)"
    grafik1.width, grafik1.height = 17, 10
    grafik1.add_data(Reference(ws_ozet, min_col=3, min_row=1, max_col=4, max_row=3), titles_from_data=True)
    kategori_ayarla(grafik1, Reference(ws_ozet, min_col=1, min_row=2, max_row=3))
    ws_db.add_chart(grafik1, "A9")

    grafik2 = PieChart()
    grafik2.title = "IBNR dağılımı"
    grafik2.width, grafik2.height = 14, 10
    grafik2.add_data(Reference(ws_ozet, min_col=4, min_row=1, max_row=3), titles_from_data=True)
    kategori_ayarla(grafik2, Reference(ws_ozet, min_col=1, min_row=2, max_row=3))
    ws_db.add_chart(grafik2, "I9")

    bolum_basligi(ws_db, "A30", "Mevcut Kümülatif ve IBNR")
    ws_db.column_dimensions['B'].width = 22
    ws_db.column_dimensions['C'].width = 22

    kopru_satir = 32
    for sutun_no, metin in enumerate(["Segment", "Mevcut Kümülatif", "IBNR"], start=1):
        hucre = ws_db.cell(row=kopru_satir, column=sutun_no, value=metin)
        hucre.font = FONT_BASLIK
        hucre.fill = RENK_LACIVERT
        hucre.alignment = Alignment(horizontal="center", vertical="center")

    kopru_veri = [("Maddî", maddi['kumulatif'], maddi['ibnr']),
                  ("Bedeni", bedeni['kumulatif'], bedeni['ibnr']),
                  ("Toplam", sonuc['toplam_kumulatif'], sonuc['toplam_ibnr'])]
    for i, (ad, kum, ibnr_deger) in enumerate(kopru_veri):
        satir_no = kopru_satir + 1 + i
        hucreler = [ws_db.cell(row=satir_no, column=1, value=ad),
                    ws_db.cell(row=satir_no, column=2, value=kum),
                    ws_db.cell(row=satir_no, column=3, value=ibnr_deger)]
        for hucre in hucreler[1:]:
            hucre.number_format = '#,##0 "TL";-#,##0 "TL"'
        for hucre in hucreler:
            hucre.border = KENAR_INCE
            hucre.alignment = Alignment(horizontal="center", vertical="center", shrink_to_fit=True)
            if i % 2 == 1:
                hucre.fill = RENK_TOPLAM
        hucreler[0].font = FONT_KALIN

    grafik3 = BarChart()
    grafik3.type = "col"
    grafik3.grouping = "stacked"
    grafik3.overlap = 100
    grafik3.title = "Mevcut kümülatif ve IBNR"
    grafik3.y_axis.title = "Tutar (TL)"
    grafik3.width, grafik3.height = 17, 10
    grafik3.add_data(Reference(ws_db, min_col=2, min_row=kopru_satir, max_col=3,
                               max_row=kopru_satir + len(kopru_veri)), titles_from_data=True)
    kategori_ayarla(grafik3, Reference(ws_db, min_col=1, min_row=kopru_satir + 1,
                                       max_row=kopru_satir + len(kopru_veri)))
    ws_db.add_chart(grafik3, "E32")

    bolum_basligi(ws_db, "A53", "IBNR Trendi (son 16 çeyrek)")
    trend = pd.DataFrame({
        'KazaDonemi': etiketler,
        'IBNR': maddi['detay']['IBNR'].values + bedeni['detay']['IBNR'].values,
    }).tail(16).reset_index(drop=True)

    trend_satir = 55
    for sutun_no, metin in enumerate(["Kaza Dönemi", "Toplam IBNR"], start=1):
        hucre = ws_db.cell(row=trend_satir, column=sutun_no, value=metin)
        hucre.font = FONT_BASLIK
        hucre.fill = RENK_LACIVERT
        hucre.alignment = Alignment(horizontal="center", vertical="center")

    for i, kayit in trend.iterrows():
        satir_no = trend_satir + 1 + i
        hucreler = [ws_db.cell(row=satir_no, column=1, value=kayit['KazaDonemi']),
                    ws_db.cell(row=satir_no, column=2, value=float(kayit['IBNR']))]
        hucreler[1].number_format = '#,##0 "TL";-#,##0 "TL"'
        for hucre in hucreler:
            hucre.border = KENAR_INCE
            hucre.alignment = Alignment(horizontal="center", vertical="center", shrink_to_fit=True)
            if i % 2 == 1:
                hucre.fill = RENK_TOPLAM

    grafik4 = LineChart()
    grafik4.title = "IBNR trendi"
    grafik4.style = 12
    grafik4.varyColors = False
    grafik4.y_axis.title = "IBNR (TL)"
    grafik4.x_axis.title = "Kaza Dönemi"
    grafik4.width, grafik4.height = 20, 10
    grafik4.add_data(Reference(ws_db, min_col=2, min_row=trend_satir,
                               max_row=trend_satir + len(trend)), titles_from_data=True)
    kategori_ayarla(grafik4, Reference(ws_db, min_col=1, min_row=trend_satir + 1,
                                       max_row=trend_satir + len(trend)))
    ws_db.add_chart(grafik4, "D55")

    not_satir = trend_satir + max(len(trend), 19) + 3
    bolum_basligi(ws_db, f"A{not_satir}", "Notlar")
    maddi_pay = (maddi['ibnr'] / sonuc['toplam_ibnr'] * 100) if sonuc['toplam_ibnr'] else 0.0
    bedeni_pay = (bedeni['ibnr'] / sonuc['toplam_ibnr'] * 100) if sonuc['toplam_ibnr'] else 0.0
    notlar = (f"{ayar['degerleme']} sonu toplam IBNR: {tl(sonuc['toplam_ibnr'])}\n"
              f"Maddî: {tl(maddi['ibnr'])} (%{maddi_pay:.1f}) | Bedeni: {tl(bedeni['ibnr'])} (%{bedeni_pay:.1f})\n"
              f"Toplam ultimate: {tl(sonuc['toplam_ultimate'])}\n"
              f"Yöntem: Chain Ladder | Üçgen: {sonuc['baslangic_donemi']}–{ayar['degerleme']} | "
              f"Faktör penceresi: {ayar['pencere']} | Kuyruk: {maddi['tail']:.4f} ({maddi['tail_not']})")
    ws_db.merge_cells(f"A{not_satir + 1}:J{not_satir + 5}")
    ws_db[f"A{not_satir + 1}"].value = notlar
    ws_db[f"A{not_satir + 1}"].font = FONT_NORMAL
    ws_db[f"A{not_satir + 1}"].alignment = Alignment(horizontal="left", vertical="top", wrap_text=True)
    ws_db[f"A{not_satir + 1}"].fill = RENK_MAVI
    for satir_no in range(not_satir + 1, not_satir + 6):
        for sutun_no in range(1, 11):
            ws_db.cell(row=satir_no, column=sutun_no).border = KENAR_KALIN

    if sonuc['senaryolar'] is not None:
        ws_senaryo = wb['Senaryolar']
        ws_senaryo.sheet_view.showGridLines = False
        baslik_satiri(ws_senaryo)
        for satir in ws_senaryo.iter_rows(min_row=2, max_row=len(sonuc['senaryolar']) + 1,
                                          min_col=2, max_col=6):
            for hucre in satir:
                if isinstance(hucre.value, (int, float)):
                    if hucre.column == 6:
                        hucre.number_format = '0.000000'
                    else:
                        hucre.number_format = '#,##0 "TL"'
                    hucre.font = FONT_NORMAL
        for hucre in ws_senaryo['G'][1:]:
            if isinstance(hucre.value, (int, float)):
                hucre.number_format = '+0.0"%";-0.0"%"'

        grafik5 = BarChart()
        grafik5.type = "col"
        grafik5.title = "Faktör penceresine göre toplam ultimate"
        grafik5.y_axis.title = "Tutar (TL)"
        grafik5.width, grafik5.height = 17, 9
        grafik5.add_data(Reference(ws_senaryo, min_col=4, min_row=1,
                                   max_row=len(sonuc['senaryolar']) + 1), titles_from_data=True)
        kategori_ayarla(grafik5, Reference(ws_senaryo, min_col=1, min_row=2,
                                           max_row=len(sonuc['senaryolar']) + 1))
        ws_senaryo.add_chart(grafik5, "I2")

    if sonuc.get('geriye_test') is not None:
        ws_test = wb['Geriye Dönük Test']
        ws_test.sheet_view.showGridLines = False
        baslik_satiri(ws_test)
        for satir in ws_test.iter_rows(min_row=2, max_row=len(sonuc['geriye_test']) + 1,
                                       min_col=4, max_col=7):
            for hucre in satir:
                if isinstance(hucre.value, (int, float)):
                    hucre.number_format = ('+0.0"%";-0.0"%"' if hucre.column == 7
                                           else '#,##0 "TL";-#,##0 "TL"')
                    hucre.font = FONT_NORMAL

    for sayfa_adi, tablo in (('Varsayımlar', varsayimlar), ('Veri Kalitesi', kalite)):
        ws_bilgi = wb[sayfa_adi]
        ws_bilgi.sheet_view.showGridLines = False
        ws_bilgi.sheet_properties.tabColor = "808080"
        baslik_satiri(ws_bilgi)
        for satir in ws_bilgi.iter_rows(min_row=2, max_row=len(tablo) + 1):
            for hucre in satir:
                hucre.font = FONT_NORMAL
                hucre.border = KENAR_INCE
                hucre.alignment = Alignment(horizontal="left", vertical="center")

    ws_kalite = wb['Veri Kalitesi']
    for satir in ws_kalite.iter_rows(min_row=2, max_row=len(kalite) + 1):
        if str(satir[2].value or "").startswith("BEKLENMEYEN"):
            for hucre in satir:
                hucre.fill = RENK_KIRMIZI
                hucre.font = FONT_KALIN

    oran_skala = ColorScaleRule(start_type='percentile', start_value=10, start_color='F8696B',
                                mid_type='percentile', mid_value=50, mid_color='FFEB84',
                                end_type='percentile', end_value=90, end_color='63BE7B')
    for anahtar, ad, _ in SEGMENTLER:
        sayfa_adi = f"{'Maddi' if anahtar == 'maddi' else 'Bedeni'} Link Ratio"
        ws_oran = wb[sayfa_adi]
        ws_oran.sheet_view.showGridLines = False
        ws_oran.freeze_panes = "B2"
        baslik_satiri(ws_oran)
        satir_sayisi = len(link_ratio[anahtar]) + 1
        sutun_sayisi = link_ratio[anahtar].shape[1] + 1
        for satir in ws_oran.iter_rows(min_row=2, max_row=satir_sayisi, min_col=2,
                                       max_col=sutun_sayisi):
            for hucre in satir:
                if isinstance(hucre.value, (int, float)):
                    hucre.number_format = '0.0000'
                    hucre.font = FONT_NORMAL
        for satir in ws_oran.iter_rows(min_row=satir_sayisi - 4, max_row=satir_sayisi,
                                       min_col=1, max_col=sutun_sayisi):
            for hucre in satir:
                hucre.fill = RENK_TOPLAM
                if hucre.column == 1 or satir[0].value == 'Kullanılan faktör':
                    hucre.font = FONT_KALIN
        ws_oran.conditional_formatting.add(
            f"B2:{get_column_letter(sutun_sayisi)}{satir_sayisi - 5}", oran_skala)

    renk_skala = ColorScaleRule(start_type='num', start_value=0, start_color='FFFFFF',
                                end_type='max', end_color='BDD7EE')
    ucgen_son_satir = periyot + 1
    ucgen_son_sutun = get_column_letter(periyot + 1)
    for sayfa_adi in ['Maddi Üçgen', 'Maddi Tamamlanmış Üçgen', 'Bedeni Üçgen', 'Bedeni Tamamlanmış Üçgen']:
        ws_ucgen = wb[sayfa_adi]
        ws_ucgen.sheet_view.showGridLines = False
        ws_ucgen.freeze_panes = "B2"
        baslik_satiri(ws_ucgen)
        for satir in ws_ucgen.iter_rows(min_row=2, max_row=ucgen_son_satir, min_col=2,
                                        max_col=periyot + 1):
            for hucre in satir:
                if isinstance(hucre.value, (int, float)) and hucre.value > 0:
                    hucre.number_format = '#,##0'
                    hucre.font = FONT_NORMAL
                elif hucre.value == 0:
                    hucre.value = ""
        ws_ucgen.conditional_formatting.add(f"B2:{ucgen_son_sutun}{ucgen_son_satir}", renk_skala)

    for sayfa_adi in ['Maddi Çeyreklik Detay', 'Bedeni Çeyreklik Detay']:
        ws_detay = wb[sayfa_adi]
        ws_detay.freeze_panes = "A2"
        baslik_satiri(ws_detay)
        for satir in ws_detay.iter_rows(min_row=2, max_row=periyot + 1, min_col=2, max_col=4):
            for hucre in satir:
                if isinstance(hucre.value, (int, float)):
                    hucre.number_format = '#,##0.00 "TL"'
                    hucre.font = FONT_NORMAL

    ws_faktor = wb.create_sheet('Gelişim Faktörleri')
    ws_faktor.append(['Gelişim Adımı', 'Maddî Faktör', 'Bedeni Faktör'])
    for j in range(son_indeks):
        ws_faktor.append([f"{j} -> {j+1}", maddi['faktorler'][j], bedeni['faktorler'][j]])
    ws_faktor.append(["Kuyruk (tail)", maddi['tail'], bedeni['tail']])
    baslik_satiri(ws_faktor)
    for satir in ws_faktor.iter_rows(min_row=2, max_row=periyot + 1, min_col=2, max_col=3):
        for hucre in satir:
            hucre.number_format = '0.000000'
            hucre.font = FONT_NORMAL

    grafik_f = LineChart()
    grafik_f.title = "Gelişim faktörleri"
    grafik_f.style = 13
    grafik_f.varyColors = False
    grafik_f.y_axis.title = "Faktör"
    grafik_f.x_axis.title = "Gelişim Dönemi"
    grafik_f.width, grafik_f.height = 18, 10
    grafik_f.add_data(Reference(ws_faktor, min_col=2, min_row=1, max_col=3, max_row=periyot),
                      titles_from_data=True)
    kategori_ayarla(grafik_f, Reference(ws_faktor, min_col=1, min_row=2, max_row=periyot))
    ws_faktor.add_chart(grafik_f, "E2")

    pbi = {'PBI_Ozet': uzun_ozet, 'PBI_Detay': uzun_detay,
           'PBI_Faktor': uzun_faktor, 'PBI_Ucgen': uzun_tumu}
    for sayfa_adi, tablo in pbi.items():
        ws_pbi = wb[sayfa_adi]
        ws_pbi.sheet_properties.tabColor = "70AD47"
        aralik = f"A1:{get_column_letter(tablo.shape[1])}{tablo.shape[0] + 1}"
        excel_tablo = Table(displayName=f"Tbl_{sayfa_adi}", ref=aralik)
        excel_tablo.tableStyleInfo = TableStyleInfo(name="TableStyleMedium2", showRowStripes=True,
                                                    showColumnStripes=False)
        ws_pbi.add_table(excel_tablo)
        baslik_satiri(ws_pbi)

    for sayfa in wb.worksheets:
        if sayfa.title == 'Dashboard':
            continue
        for sutun in sayfa.columns:
            en_uzun = max(len(str(hucre.value or '')) for hucre in sutun)
            sayfa.column_dimensions[get_column_letter(sutun[0].column)].width = max(en_uzun + 4, 15)

    wb.save(cikis)
    return cikis


def cikis_yolu(ayar):
    ozel = os.environ.get("IBNR_CIKIS_DOSYASI")
    if ozel:
        return ozel
    klasor = os.path.dirname(os.path.abspath(ayar['giris']))
    return os.path.join(klasor, f"IBNR_Rapor_{ayar['degerleme']}.xlsx")


def konsola_yaz(sonuc, ayar):
    print(f"\nÜçgen: {sonuc['baslangic_donemi']} — {ayar['degerleme']} ({sonuc['periyot']} çeyrek)"
          f" | Veri diyagonali: {sonuc['veri_son_donemi']}")
    print(f"Branş {ayar['brans']} | Havuz {ayar['havuz']} | Faktör {ayar['pencere']} | "
          f"Kuyruk {sonuc['maddi']['tail']:.4f} ({sonuc['maddi']['tail_not']}) | "
          f"Sovtaj {tam_sayi(ayar['sovtaj']):+d}")
    print("-" * 70)
    print(sonuc['ozet'])
    if sonuc['senaryolar'] is not None:
        print("-" * 70)
        print(sonuc['senaryolar'].to_string(index=False))
    if sonuc.get('geriye_test') is not None:
        print("-" * 70)
        print("Geriye dönük test")
        print(sonuc['geriye_test'].to_string(index=False))
    print("-" * 70)


def gui_var():
    if not INTERAKTIF:
        return False
    try:
        import tkinter
        kok = tkinter.Tk()
        kok.destroy()
        return True
    except Exception:
        return False


def dosya_sec(varsayilan):
    if os.path.isfile(varsayilan):
        return varsayilan
    if not gui_var():
        raise SystemExit(f"Dosya bulunamadı: {varsayilan}")
    from tkinter import Tk, filedialog
    kok = Tk()
    kok.withdraw()
    yol = filedialog.askopenfilename(title="Hasar verisi dosyasını seçin",
                                     filetypes=[("Excel", "*.xlsx *.xlsm *.xls"), ("Tümü", "*.*")])
    kok.destroy()
    if not yol:
        raise SystemExit("Dosya seçilmedi.")
    return yol


def donem_secenekleri(meta, ileri=8):
    return [kod_donem(k) for k in range(meta['veri_ilk'], meta['veri_son'] + ileri + 1)]


def arayuz(df, ayar, meta):
    import tkinter as tk
    from tkinter import ttk, messagebox

    durum_kaydi = {'sonuc': None, 'cikis': None, 'override': None}

    pencere = tk.Tk()
    pencere.title("IBNR Hesaplama")
    pencere.resizable(False, False)

    kap = ttk.Frame(pencere, padding=16)
    kap.grid(row=0, column=0, sticky="nsew")
    kap.columnconfigure(1, weight=1)

    bilgi = (f"Dosya: {os.path.basename(meta['yol'])}\n"
             f"Veri aralığı: {kod_donem(meta['veri_ilk'])} — {kod_donem(meta['veri_son'])} "
             f"({meta['veri_son'] - meta['veri_ilk'] + 1} çeyrek)")
    ttk.Label(kap, text=bilgi, justify="left").grid(row=0, column=0, columnspan=2,
                                                    sticky="w", pady=(0, 12))

    donemler = donem_secenekleri(meta)
    azami = donem_kodu(ayar['degerleme']) - meta['veri_ilk'] + 1
    ceyrek_secenekleri = [str(x) for x in [4, 8, 12, 16, 20, 24, 28, 32, 36, 40, 44, 48] if x <= azami]
    if str(azami) not in ceyrek_secenekleri:
        ceyrek_secenekleri.append(str(azami))

    alan = {}

    etiketler_gerec = {}

    def satir(indeks, etiket, gerec):
        etiket_gerec = ttk.Label(kap, text=etiket)
        etiket_gerec.grid(row=indeks, column=0, sticky="w", pady=4, padx=(0, 12))
        gerec.grid(row=indeks, column=1, sticky="ew", pady=4)
        etiketler_gerec[indeks] = etiket_gerec
        return etiket_gerec

    alan['degerleme'] = ttk.Combobox(kap, width=22, values=donemler)
    alan['degerleme'].set(str(ayar['degerleme']))
    satir(1, "Değerleme dönemi", alan['degerleme'])

    alan['brans'] = ttk.Combobox(kap, width=22, values=meta['brans_secenekleri'])
    alan['brans'].set(sayi_metni(ayar['brans']))
    satir(2, "Branş kodu", alan['brans'])

    alan['havuz'] = ttk.Combobox(kap, width=22, values=meta['havuz_secenekleri'])
    alan['havuz'].set(sayi_metni(ayar['havuz']))
    satir(3, "Havuz kodu", alan['havuz'])

    alan['ceyrek'] = ttk.Combobox(kap, width=22, values=ceyrek_secenekleri)
    alan['ceyrek'].set(sayi_metni(min(tam_sayi(ayar['ceyrek']), azami)))
    ceyrek_etiketi = satir(4, f"Üçgen kaç çeyrek (azami {azami})", alan['ceyrek'])

    alan['pencere'] = ttk.Combobox(kap, width=22, state="readonly", values=PENCERE_SECENEKLERI)
    alan['pencere'].set(str(ayar['pencere']))
    satir(5, "Faktör penceresi", alan['pencere'])

    alan['tail'] = ttk.Entry(kap, width=24)
    alan['tail'].insert(0, str(ayar['tail']))
    satir(6, "Kuyruk (tail) faktörü", alan['tail'])

    tail_oto = tk.BooleanVar(value=bool(ayar['tail_otomatik']))

    def tail_durum():
        alan['tail'].configure(state="disabled" if tail_oto.get() else "normal")

    ttk.Checkbutton(kap, text="Kuyruk faktörünü verilerden tahmin et",
                    variable=tail_oto, command=tail_durum).grid(row=7, column=1, sticky="w")

    alan['sovtaj'] = ttk.Combobox(kap, width=22, state="readonly", values=['1', '-1'])
    alan['sovtaj'].set(sayi_metni(ayar['sovtaj']))
    satir(8, "Sovtaj işareti (1 artırır, -1 azaltır)", alan['sovtaj'])

    senaryo_var = tk.BooleanVar(value=bool(ayar['senaryolar']))
    ttk.Checkbutton(kap, text="Üç faktör penceresini de hesapla ve karşılaştır",
                    variable=senaryo_var).grid(row=9, column=1, sticky="w", pady=(4, 0))

    test_var = tk.BooleanVar(value=bool(ayar.get('geriye_test', True)))
    ttk.Checkbutton(kap, text="Geriye dönük test yap (4/8/12 çeyrek önce)",
                    variable=test_var).grid(row=10, column=1, sticky="w")

    durum_etiketi = ttk.Label(kap, text="", foreground="#1F4E78")
    durum_etiketi.grid(row=11, column=0, columnspan=2, sticky="w", pady=(12, 0))

    def durum(metin):
        durum_etiketi.configure(text=metin)
        pencere.update_idletasks()
        pencere.update()

    def formu_oku():
        degerleme = kod_donem(donem_kodu(alan['degerleme'].get()))
        ceyrek = tam_sayi(alan['ceyrek'].get())
        if ceyrek < 2:
            raise ValueError("Üçgen en az 2 çeyrek olmalı")
        yeni = dict(ayar)
        yeni.update({
            'degerleme': degerleme,
            'degerleme_kodu': donem_kodu(degerleme),
            'brans': tam_sayi(alan['brans'].get()),
            'havuz': tam_sayi(alan['havuz'].get()),
            'ceyrek': ceyrek,
            'pencere': pencere_dogrula(alan['pencere'].get()),
            'tail': 1.0 if tail_oto.get() else tail_dogrula(alan['tail'].get()),
            'tail_otomatik': bool(tail_oto.get()),
            'sovtaj': isaret_dogrula(alan['sovtaj'].get()),
            'senaryolar': bool(senaryo_var.get()),
            'geriye_test': bool(test_var.get()),
        })
        return yeni

    def donem_degisti(olay=None):
        try:
            kod = donem_kodu(alan['degerleme'].get())
        except ValueError:
            return
        yeni_azami = max(kod - meta['veri_ilk'] + 1, 2)
        secenekler = [str(x) for x in range(4, yeni_azami + 1, 4)]
        if str(yeni_azami) not in secenekler:
            secenekler.append(str(yeni_azami))
        alan['ceyrek'].configure(values=secenekler)
        ceyrek_etiketi.configure(text=f"Üçgen kaç çeyrek (azami {yeni_azami})")
        try:
            if tam_sayi(alan['ceyrek'].get()) > yeni_azami:
                alan['ceyrek'].set(str(yeni_azami))
        except ValueError:
            alan['ceyrek'].set(str(yeni_azami))

    def calistir():
        try:
            guncel = formu_oku()
            uyarilar = ayari_dogrula(guncel, meta)
        except (ValueError, TypeError) as hata:
            messagebox.showerror("Hatalı giriş", str(hata))
            return
        if uyarilar:
            messagebox.showinfo("Ayar düzeltildi", "\n".join(uyarilar))
            alan['ceyrek'].set(sayi_metni(guncel['ceyrek']))
            alan['degerleme'].set(guncel['degerleme'])
        try:
            durum("Hesaplanıyor...")
            sonuc = hesapla(df, guncel, meta, durum=durum)
            cikis = cikis_yolu(guncel)
            excel_yaz(sonuc, guncel, meta, cikis, durum=durum)
        except Exception as hata:
            durum("")
            messagebox.showerror("Hesaplama hatası", str(hata))
            return
        ayarlari_yaz(guncel)
        konsola_yaz(sonuc, guncel)
        durum_kaydi.update({'sonuc': sonuc, 'cikis': cikis, 'ayar': guncel})
        sonuc_ekrani(pencere, df, meta, durum_kaydi)

    dugmeler = ttk.Frame(kap)
    dugmeler.grid(row=12, column=0, columnspan=2, pady=(16, 0), sticky="e")
    ttk.Button(dugmeler, text="Kapat", command=pencere.destroy).grid(row=0, column=0, padx=(0, 8))
    ttk.Button(dugmeler, text="Hesapla", command=calistir).grid(row=0, column=1)

    alan['degerleme'].bind("<<ComboboxSelected>>", donem_degisti)
    alan['degerleme'].bind("<FocusOut>", donem_degisti)
    pencere.bind("<Return>", lambda olay: calistir())
    tail_durum()
    pencere.mainloop()
    return durum_kaydi


def sonuc_ekrani(ana_pencere, df, meta, kayit):
    import tkinter as tk
    from tkinter import ttk, messagebox

    sonuc, ayar, cikis = kayit['sonuc'], kayit['ayar'], kayit['cikis']

    ekran = tk.Toplevel(ana_pencere)
    ekran.title("Sonuç")
    ekran.resizable(False, False)
    kap = ttk.Frame(ekran, padding=16)
    kap.grid(row=0, column=0)

    baslik = (f"{ayar['degerleme']} değerlemesi | {sonuc['baslangic_donemi']}–{ayar['degerleme']} "
              f"({sonuc['periyot']} çeyrek) | Branş {ayar['brans']}")
    ttk.Label(kap, text=baslik, font=("Segoe UI", 10, "bold")).grid(
        row=0, column=0, columnspan=3, sticky="w", pady=(0, 12))

    tablo = [
        ("", "Mevcut Kümülatif", "IBNR", "Ultimate"),
        ("Maddî", sonuc['maddi']['kumulatif'], sonuc['maddi']['ibnr'], sonuc['maddi']['ultimate']),
        ("Bedeni", sonuc['bedeni']['kumulatif'], sonuc['bedeni']['ibnr'], sonuc['bedeni']['ultimate']),
        ("Toplam", sonuc['toplam_kumulatif'], sonuc['toplam_ibnr'], sonuc['toplam_ultimate']),
    ]
    for r, satir in enumerate(tablo):
        for c, deger in enumerate(satir):
            metin = deger if isinstance(deger, str) else tl(deger)
            kalin = (r == 0 or r == 3)
            ttk.Label(kap, text=metin, font=("Segoe UI", 9, "bold" if kalin else "normal")).grid(
                row=1 + r, column=c, sticky="e" if c else "w", padx=8, pady=2)

    satir_no = 6
    if sonuc['senaryolar'] is not None:
        ttk.Separator(kap, orient="horizontal").grid(row=satir_no, column=0, columnspan=4,
                                                     sticky="ew", pady=10)
        satir_no += 1
        ttk.Label(kap, text="Faktör penceresi karşılaştırması",
                  font=("Segoe UI", 9, "bold")).grid(row=satir_no, column=0, columnspan=4, sticky="w")
        satir_no += 1
        for _, kayit_satiri in sonuc['senaryolar'].iterrows():
            fark = float(kayit_satiri[FARK_KOLONU])
            metin = (f"{kayit_satiri['Faktör Penceresi']:8}  "
                     f"ultimate {tl(kayit_satiri['Toplam Ultimate'])}   "
                     f"IBNR {tl(kayit_satiri['Toplam IBNR'])}   ({fark:+.1f}%)")
            ttk.Label(kap, text=metin).grid(row=satir_no, column=0, columnspan=4, sticky="w")
            satir_no += 1

    if sonuc.get('geriye_test') is not None and len(sonuc['geriye_test']):
        ttk.Separator(kap, orient="horizontal").grid(row=satir_no, column=0, columnspan=4,
                                                     sticky="ew", pady=10)
        satir_no += 1
        ttk.Label(kap, text="Geriye dönük test (o günkü tahmin, bugüne göre)",
                  font=("Segoe UI", 9, "bold")).grid(row=satir_no, column=0, columnspan=4, sticky="w")
        satir_no += 1
        for _, test_satiri in sonuc['geriye_test'].iterrows():
            metin = (f"{int(test_satiri['Gecikme (çeyrek)'])} çeyrek önce "
                     f"({test_satiri['Değerleme Dönemi']}):  sapma "
                     f"{float(test_satiri['Sapma %']):+.1f}%")
            ttk.Label(kap, text=metin).grid(row=satir_no, column=0, columnspan=4, sticky="w")
            satir_no += 1

    ttk.Separator(kap, orient="horizontal").grid(row=satir_no, column=0, columnspan=4,
                                                 sticky="ew", pady=10)
    satir_no += 1
    ttk.Label(kap, text=f"Rapor: {cikis}", foreground="#595959").grid(
        row=satir_no, column=0, columnspan=4, sticky="w")
    satir_no += 1

    def yeniden_hesapla(override):
        try:
            yeni_sonuc = hesapla(df, ayar, meta, override=override)
            excel_yaz(yeni_sonuc, ayar, meta, cikis)
        except Exception as hata:
            messagebox.showerror("Hesaplama hatası", str(hata))
            return
        kayit.update({'sonuc': yeni_sonuc, 'override': override})
        konsola_yaz(yeni_sonuc, ayar)
        ekran.destroy()
        sonuc_ekrani(ana_pencere, df, meta, kayit)

    dugmeler = ttk.Frame(kap)
    dugmeler.grid(row=satir_no, column=0, columnspan=4, pady=(14, 0), sticky="e")
    ttk.Button(dugmeler, text="Excel'i aç", command=lambda: dosya_ac(cikis)).grid(row=0, column=0, padx=4)
    ttk.Button(dugmeler, text="Faktörleri düzenle",
               command=lambda: faktor_ekrani(ekran, sonuc, yeniden_hesapla)).grid(row=0, column=1, padx=4)
    ttk.Button(dugmeler, text="Kapat", command=ana_pencere.destroy).grid(row=0, column=2, padx=4)


def faktor_ekrani(ust_pencere, sonuc, kaydet):
    import tkinter as tk
    from tkinter import ttk, messagebox

    ekran = tk.Toplevel(ust_pencere)
    ekran.title("Gelişim faktörleri")
    ekran.geometry("460x520")

    ttk.Label(ekran, text="Hesaplanan faktörleri değiştirip kaydedebilirsiniz.",
              padding=10).grid(row=0, column=0, sticky="w")

    tuval = tk.Canvas(ekran, width=430, height=400, highlightthickness=0)
    kaydirma = ttk.Scrollbar(ekran, orient="vertical", command=tuval.yview)
    ic_cerceve = ttk.Frame(tuval)
    ic_cerceve.bind("<Configure>", lambda olay: tuval.configure(scrollregion=tuval.bbox("all")))
    tuval.create_window((0, 0), window=ic_cerceve, anchor="nw")
    tuval.configure(yscrollcommand=kaydirma.set)
    tuval.grid(row=1, column=0, padx=(10, 0), sticky="nsew")
    kaydirma.grid(row=1, column=1, sticky="ns")

    for sutun_no, baslik in enumerate(["Adım", "Maddî", "Bedeni"]):
        ttk.Label(ic_cerceve, text=baslik, font=("Segoe UI", 9, "bold")).grid(
            row=0, column=sutun_no, padx=6, pady=4)

    girisler = {'maddi': [], 'bedeni': []}
    son_indeks = sonuc['son_indeks']
    for j in range(son_indeks):
        ttk.Label(ic_cerceve, text=f"{j} → {j+1}").grid(row=j + 1, column=0, padx=6, pady=2, sticky="w")
        for sutun_no, anahtar in enumerate(['maddi', 'bedeni'], start=1):
            giris = ttk.Entry(ic_cerceve, width=14)
            giris.insert(0, f"{sonuc[anahtar]['faktorler'][j]:.6f}")
            giris.grid(row=j + 1, column=sutun_no, padx=6, pady=2)
            girisler[anahtar].append(giris)

    kuyruk_satir = son_indeks + 1
    ttk.Label(ic_cerceve, text="Kuyruk", font=("Segoe UI", 9, "bold")).grid(
        row=kuyruk_satir, column=0, padx=6, pady=(8, 2), sticky="w")
    for sutun_no, anahtar in enumerate(['maddi', 'bedeni'], start=1):
        giris = ttk.Entry(ic_cerceve, width=14)
        giris.insert(0, f"{sonuc[anahtar]['tail']:.6f}")
        giris.grid(row=kuyruk_satir, column=sutun_no, padx=6, pady=(8, 2))
        girisler[anahtar].append(giris)

    def uygula():
        override = {}
        try:
            for anahtar in ('maddi', 'bedeni'):
                degerler = [float(str(giris.get()).replace(',', '.')) for giris in girisler[anahtar]]
                if any(d <= 0 for d in degerler):
                    raise ValueError("Faktörler pozitif olmalı")
                override[anahtar] = degerler
        except ValueError as hata:
            messagebox.showerror("Hatalı faktör", str(hata))
            return
        ekran.destroy()
        kaydet(override)

    dugmeler = ttk.Frame(ekran, padding=10)
    dugmeler.grid(row=2, column=0, columnspan=2, sticky="e")
    ttk.Button(dugmeler, text="Vazgeç", command=ekran.destroy).grid(row=0, column=0, padx=4)
    ttk.Button(dugmeler, text="Kaydet ve yeniden hesapla", command=uygula).grid(row=0, column=1, padx=4)


def main():
    ayar = ayarlari_oku()
    ayar['giris'] = dosya_sec(ayar['giris'])
    df, meta = veri_oku(ayar['giris'])

    if not ayar.get('degerleme'):
        ayar['degerleme'] = kod_donem(meta['veri_son'])
    try:
        donem_kodu(ayar['degerleme'])
    except ValueError:
        ayar['degerleme'] = kod_donem(meta['veri_son'])
    ayar['degerleme_kodu'] = donem_kodu(ayar['degerleme'])

    if gui_var():
        arayuz(df, ayar, meta)
        return

    for uyari in ayari_dogrula(ayar, meta):
        print(f"  ! {uyari}")
    sonuc = hesapla(df, ayar, meta)
    cikis = cikis_yolu(ayar)
    excel_yaz(sonuc, ayar, meta, cikis)
    konsola_yaz(sonuc, ayar)
    print(f"Çıktı: {cikis}")


if __name__ == "__main__":
    main()